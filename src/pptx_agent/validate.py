"""v1 acceptance check (spec Phasing table): does the generated .pptx actually
hold editable text (not just picture layers), and does every slide still have
at least one non-text visual element after translation?"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation


@dataclass
class SlideReport:
    index: int
    text_shape_count: int
    non_text_shape_count: int
    picture_count: int
    chart_count: int = 0
    problems: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problems


def validate_pptx(run_dir: Path) -> tuple[bool, list[SlideReport]]:
    pptx_path = run_dir / "deck.pptx"
    if not pptx_path.exists():
        raise FileNotFoundError(f"{pptx_path} not found — run the Node extract/translate stage first")

    prs = Presentation(str(pptx_path))
    reports: list[SlideReport] = []

    expected_slide_count = None
    content_path = run_dir / "content.json"
    if content_path.exists():
        expected_slide_count = len(json.loads(content_path.read_text(encoding="utf-8"))["slides"])

    for i, slide in enumerate(prs.slides):
        text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        picture_shapes = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        chart_shapes = [s for s in slide.shapes if getattr(s, "has_chart", False)]
        non_text_shapes = [s for s in slide.shapes if not (s.has_text_frame and s.text_frame.text.strip())]

        report = SlideReport(
            index=i,
            text_shape_count=len(text_shapes),
            non_text_shape_count=len(non_text_shapes),
            picture_count=len(picture_shapes),
            chart_count=len(chart_shapes),
        )
        if not text_shapes:
            report.problems.append("no editable text shapes found (deck may have collapsed to an image)")
        if not non_text_shapes:
            report.problems.append("no non-text visual element found (fails 'not plain title+bullets' MVP criterion)")
        reports.append(report)

    if expected_slide_count is not None and len(prs.slides) != expected_slide_count:
        reports.append(
            SlideReport(
                index=-1,
                text_shape_count=0,
                non_text_shape_count=0,
                picture_count=0,
                problems=[f"slide count mismatch: content.json has {expected_slide_count}, pptx has {len(prs.slides)}"],
            )
        )

    return all(r.ok for r in reports), reports


def main(argv: list[str] | None = None) -> int:
    import sys

    args = argv if argv is not None else sys.argv[1:]
    if not args:
        print("Usage: python -m pptx_agent.validate <run_dir>", file=sys.stderr)
        return 2
    run_dir = Path(args[0])
    ok, reports = validate_pptx(run_dir)
    for r in reports:
        status = "OK" if r.ok else "FAIL"
        label = f"slide {r.index}" if r.index >= 0 else "deck"
        print(f"[{status}] {label}: text_shapes={r.text_shape_count} non_text_shapes={r.non_text_shape_count} pictures={r.picture_count} charts={r.chart_count}")
        for p in r.problems:
            print(f"       - {p}")
    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
